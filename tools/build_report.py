"""Build the laboratory report from the supplied DOCX reference."""

from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path(
    "/Users/denyszastavnyi/Downloads/Telegram Desktop/Приклад_звіту.docx"
)
EXPECTED_REFERENCE_HASH = (
    "1b16de3fea64f57b4071549ff5e567d1ec85845283e4e3c4b457a8518c3bfa57"
)
OUTPUT = ROOT / "docs" / "Lab_01_Report_Variant_7.docx"
SCREENSHOTS = ROOT / "docs" / "screenshots"


def set_run_font(run, name: str, size: float, *, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:cs"), name)


def clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")

    for style_name, size, color in (
        ("Heading 1", 14, "366091"),
        ("Heading 2", 13, "4F81BD"),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:cs"), "Calibri")


def add_cover_paragraph(
    document: Document,
    text: str,
    *,
    size: float = 14,
    bold: bool = False,
    before: float = 0,
    after: float = 0,
    left_indent: float | None = None,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
        if left_indent is not None
        else WD_ALIGN_PARAGRAPH.CENTER
    )
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    if left_indent is not None:
        paragraph.paragraph_format.left_indent = Inches(left_indent)
    run = paragraph.add_run(text)
    set_run_font(run, "Times New Roman", size, bold=bold)


def add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(document: Document, text: str, *, indent: bool = True):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Inches(0.49 if indent else 0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(text)
    set_run_font(run, "Times New Roman", 12)
    return paragraph


def add_code(document: Document, text: str, *, size: float = 8.5):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.first_line_indent = Inches(0)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text.rstrip())
    set_run_font(run, "Consolas", size)
    return paragraph


def add_picture(
    document: Document,
    filename: str,
    caption: str,
    alt_text: str,
    *,
    width: float = 6.45,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    shape = paragraph.add_run().add_picture(
        str(SCREENSHOTS / filename), width=Inches(width)
    )
    shape._inline.docPr.set("descr", alt_text)

    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(4)
    caption_paragraph.paragraph_format.space_after = Pt(8)
    caption_paragraph.paragraph_format.line_spacing = 1.0
    caption_paragraph.paragraph_format.keep_together = True
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, "Times New Roman", 10)
    caption_run.italic = True


def add_question_answer(
    document: Document,
    number: int,
    question: str,
    answer: str,
) -> None:
    question_paragraph = document.add_paragraph()
    question_paragraph.paragraph_format.keep_with_next = True
    question_paragraph.paragraph_format.space_before = Pt(4)
    question_paragraph.paragraph_format.space_after = Pt(2)
    question_run = question_paragraph.add_run(f"{number}. {question}")
    set_run_font(question_run, "Times New Roman", 12, bold=True)

    answer_paragraph = document.add_paragraph()
    answer_paragraph.paragraph_format.left_indent = Inches(0.2)
    answer_paragraph.paragraph_format.space_after = Pt(6)
    answer_run = answer_paragraph.add_run(answer)
    set_run_font(answer_run, "Times New Roman", 12)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def build_report() -> None:
    actual_hash = hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
    if actual_hash != EXPECTED_REFERENCE_HASH:
        raise RuntimeError("Reference document changed; distillation must be repeated.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    document = Document(OUTPUT)
    clear_document_body(document)
    configure_styles(document)

    footer = document.sections[0].footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.clear()
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run("Лабораторна робота №1 | Варіант №7")
    set_run_font(footer_run, "Times New Roman", 9)

    add_cover_paragraph(
        document,
        "ЛЬВІВСЬКИЙ НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ\nІМЕНІ ІВАНА ФРАНКА",
        bold=True,
        after=86,
    )
    add_cover_paragraph(
        document,
        "ЗВІТ\nДО ЛАБОРАТОРНОЇ РОБОТИ №1",
        size=16,
        bold=True,
        after=16,
    )
    add_cover_paragraph(
        document,
        "«Організація професійного Python-проєкту та Git»",
        bold=True,
        after=18,
    )
    add_cover_paragraph(
        document,
        "з курсу «Професійний Python»\n"
        "Варіант №7 — «Система оцінювання студентів»",
        after=58,
    )
    add_cover_paragraph(
        document,
        "Виконав(ла): ______________________________",
        after=10,
        left_indent=3.15,
    )
    add_cover_paragraph(
        document,
        "Група: ___________________________________",
        after=10,
        left_indent=3.15,
    )
    add_cover_paragraph(
        document,
        "Перевірив(ла): ____________________________",
        after=10,
        left_indent=3.15,
    )
    add_cover_paragraph(
        document,
        "Дата: ____________________________________",
        after=58,
        left_indent=3.15,
    )
    add_cover_paragraph(document, "Львів — 2026")
    document.add_page_break()

    add_heading(document, "1. Мета роботи")
    add_body(
        document,
        "Ознайомитися з принципами організації професійного Python-проєкту, "
        "створити ізольоване середовище, застосувати src-layout, модулі та "
        "пакети Python, налаштувати pyproject.toml, Git і документацію. "
        "Практично реалізувати консольний застосунок за варіантом №7 та "
        "перевірити його автоматизованими тестами.",
    )

    add_heading(document, "2. Індивідуальний варіант і постановка задачі")
    add_body(
        document,
        "Варіант №7 — «Система оцінювання студентів». Для кожного запису "
        "зберігаються ПІБ студента, дисципліна та набір оцінок. Програма "
        "обчислює середній бал, мінімальну й максимальну оцінки, визначає "
        "рівень успішності, сортує студентів і знаходить записи із середнім "
        "балом вище заданого.",
    )
    add_body(
        document,
        "Проєкт має src-layout, окремі модулі моделі даних, business logic, "
        "конфігурації, винятків та application entry point. Реалізовано "
        "type hints, інтерактивне меню, введення з клавіатури, перевірку "
        "даних, пошук за кількома критеріями, статистику, власний виняток, "
        "dataclass і автоматизовані тести.",
    )

    add_heading(document, "3. Структура проєкту")
    add_code(
        document,
        "student-grading-system/\n"
        "├── .gitignore\n"
        "├── pyproject.toml\n"
        "├── README.md\n"
        "├── docs/\n"
        "│   └── Lab_01_Report_Variant_7.docx\n"
        "├── src/\n"
        "│   └── student_grading/\n"
        "│       ├── __init__.py\n"
        "│       ├── __main__.py\n"
        "│       ├── config.py\n"
        "│       ├── exceptions.py\n"
        "│       ├── main.py\n"
        "│       ├── models.py\n"
        "│       └── services.py\n"
        "└── tests/\n"
        "    ├── test_main.py\n"
        "    ├── test_models.py\n"
        "    └── test_services.py",
    )
    add_body(
        document,
        "Каталог src/student_grading містить встановлюваний пакет. Каталог "
        "tests відокремлює автоматизовані перевірки від production code. "
        "Файли конфігурації й документації розташовані в корені репозиторію.",
    )

    add_heading(document, "4. Опис створених модулів")
    add_heading(document, "4.1. models.py", 2)
    add_body(
        document,
        "Модель StudentRecord реалізовано як незмінний dataclass зі slots. "
        "Конструктор очищає текстові поля, перетворює оцінки на tuple чисел "
        "і відхиляє порожні, нескінченні або позадіапазонні значення. "
        "Властивості моделі повертають середній бал, мінімум і максимум.",
    )
    add_code(
        document,
        "@dataclass(frozen=True, slots=True)\n"
        "class StudentRecord:\n"
        "    full_name: str\n"
        "    subject: str\n"
        "    grades: tuple[float, ...]\n\n"
        "    @property\n"
        "    def average_grade(self) -> float:\n"
        "        return sum(self.grades) / len(self.grades)\n\n"
        "    @property\n"
        "    def minimum_grade(self) -> float:\n"
        "        return min(self.grades)\n\n"
        "    @property\n"
        "    def maximum_grade(self) -> float:\n"
        "        return max(self.grades)",
    )

    add_heading(document, "4.2. services.py", 2)
    add_body(
        document,
        "Модуль містить business logic без залежності від консолі. Функції "
        "calculate_average(), find_min_max_grades(), determine_performance(), "
        "sort_students() і find_students_above_average() безпосередньо "
        "реалізують вимоги варіанта. Додатково реалізовано багатокритеріальний "
        "пошук і загальну статистику.",
    )
    add_code(
        document,
        "def determine_performance(student: StudentRecord) -> PerformanceLevel:\n"
        "    average = calculate_average(student)\n"
        "    if average >= EXCELLENT_GRADE:\n"
        "        return PerformanceLevel.EXCELLENT\n"
        "    if average >= GOOD_GRADE:\n"
        "        return PerformanceLevel.GOOD\n"
        "    if average >= PASSING_GRADE:\n"
        "        return PerformanceLevel.SATISFACTORY\n"
        "    return PerformanceLevel.FAILING\n\n"
        "def find_students_above_average(\n"
        "    students: Iterable[StudentRecord],\n"
        "    threshold: float,\n"
        ") -> list[StudentRecord]:\n"
        "    normalized_threshold = _validate_threshold(threshold)\n"
        "    return [\n"
        "        student for student in students\n"
        "        if student.average_grade > normalized_threshold\n"
        "    ]",
    )

    add_heading(document, "4.3. main.py", 2)
    add_body(
        document,
        "Модуль створює демонстраційні записи, форматує результати, приймає "
        "аргументи командного рядка й запускає menu-driven interface. "
        "Звичайний запуск виконує детерміновану демонстрацію та не зависає в "
        "очікуванні вводу. Прапорець --interactive відкриває меню.",
    )
    add_code(
        document,
        "def main(argv: Sequence[str] | None = None) -> None:\n"
        "    arguments = build_parser().parse_args(argv)\n"
        "    if arguments.interactive:\n"
        "        run_menu()\n"
        "        return\n"
        "    run_demo(arguments.threshold)\n\n"
        "if __name__ == \"__main__\":\n"
        "    main()",
    )

    add_heading(document, "4.4. config.py і exceptions.py", 2)
    add_body(
        document,
        "config.py зберігає межі 100-бальної шкали, прохідний бал і типовий "
        "поріг пошуку. exceptions.py визначає GradeValidationError, через який "
        "модель і сервісний шар однаково повідомляють про некоректні дані.",
    )
    add_code(
        document,
        "MIN_GRADE = 0.0\n"
        "MAX_GRADE = 100.0\n"
        "PASSING_GRADE = 60.0\n"
        "GOOD_GRADE = 75.0\n"
        "EXCELLENT_GRADE = 90.0\n\n"
        "class GradeValidationError(ValueError):\n"
        "    \"\"\"Raised when a student record or grade value is invalid.\"\"\"",
    )

    add_heading(document, "4.5. Автоматизовані тести", 2)
    add_body(
        document,
        "Три тестові модулі перевіряють модель, усі обов'язкові функції "
        "варіанта, граничні значення шкали, помилки введення, сортування без "
        "зміни вхідного списку, багатокритеріальний пошук, статистику, "
        "демонстраційний режим та всі команди меню.",
    )

    add_heading(document, "5. Файл pyproject.toml")
    add_body(
        document,
        "Файл визначає build system setuptools, метадані пакета, Python 3.11+, "
        "editable installation, console script, тестові залежності та "
        "налаштування pytest і coverage. Поріг покриття встановлено на 90%.",
    )
    add_code(
        document,
        "[build-system]\n"
        "requires = [\"setuptools>=70\"]\n"
        "build-backend = \"setuptools.build_meta\"\n\n"
        "[project]\n"
        "name = \"student-grading-system\"\n"
        "version = \"0.1.0\"\n"
        "requires-python = \">=3.11\"\n\n"
        "[project.optional-dependencies]\n"
        "test = [\"pytest>=8.0\", \"pytest-cov>=5.0\"]\n\n"
        "[project.scripts]\n"
        "student-grading = \"student_grading.main:main\"\n\n"
        "[tool.pytest.ini_options]\n"
        "pythonpath = [\"src\"]\n"
        "testpaths = [\"tests\"]\n\n"
        "[tool.coverage.report]\n"
        "fail_under = 90",
    )

    add_heading(document, "6. Файл .gitignore")
    gitignore_intro = add_body(
        document,
        "Git не відстежує virtual environment, Python caches, результати "
        "покриття, IDE configuration, build artifacts, .env і системні файли.",
    )
    gitignore_intro.paragraph_format.keep_with_next = True
    add_code(
        document,
        ".venv/\n"
        "__pycache__/\n"
        "*.py[cod]\n"
        ".pytest_cache/\n"
        ".coverage\n"
        "htmlcov/\n"
        ".mypy_cache/\n"
        ".ruff_cache/\n"
        ".idea/\n"
        ".vscode/\n"
        "dist/\n"
        "build/\n"
        "*.egg-info/\n"
        ".env\n"
        ".DS_Store",
    )

    add_heading(document, "7. Фрагмент README.md")
    add_code(
        document,
        "# Student Grading System\n\n"
        "Лабораторна робота №1 з курсу «Професійний Python».\n"
        "Варіант №7 — «Система оцінювання студентів».\n\n"
        "## Встановлення\n"
        "python3 -m venv .venv\n"
        "source .venv/bin/activate\n"
        "python -m pip install -e \".[test]\"\n\n"
        "## Запуск\n"
        "python -m student_grading\n"
        "python -m student_grading --interactive\n\n"
        "## Тести\n"
        "python -m pytest --cov=student_grading --cov-report=term-missing",
    )

    add_heading(document, "8. Створення virtual environment та запуск")
    add_code(
        document,
        "python3 -m venv .venv\n"
        "source .venv/bin/activate\n"
        "python -m pip install -e \".[test]\"\n"
        "python -m student_grading --threshold 80\n"
        "python -m student_grading --interactive",
    )
    add_body(
        document,
        "Пакет встановлено в editable mode. Команди python -m student_grading "
        "та student-grading успішно запускають application entry point.",
    )

    add_heading(document, "9. Результат роботи програми")
    add_body(
        document,
        "Під час контрольного запуску програма вивела всі записи, рейтинг, "
        "студентів із середнім балом вище 80 і загальну статистику. Значення "
        "середнього бала, мінімуму, максимуму й статусу розраховані з фактичних "
        "демонстраційних даних.",
    )
    add_picture(
        document,
        "application_output_1.png",
        "Рисунок 1 — Виведення записів, рейтингу та результату пошуку",
        "Terminal output with student records, ranking, and students above threshold",
    )
    add_picture(
        document,
        "application_output_2.png",
        "Рисунок 2 — Загальна статистика",
        "Terminal output with overall grade statistics",
    )

    add_heading(document, "10. Перевірка та експериментальна частина")
    add_body(
        document,
        "Тестовий набір містить 38 перевірок. Він охоплює валідні сценарії, "
        "порожні значення, оцінки за межами 0–100, нескінченні значення, "
        "порогові рівні 60, 75 і 90, строгий пошук вище порога, case-insensitive "
        "пошук, усі команди меню та обидва режими запуску. Усі тести пройшли; "
        "coverage становить 99,6%, що перевищує налаштований поріг 90%.",
    )
    add_picture(
        document,
        "test_results.png",
        "Рисунок 3 — Тести та coverage",
        "Pytest output showing 38 passing tests and 99.6 percent coverage",
        width=5.5,
    )

    add_heading(document, "11. Робота з Git")
    add_body(
        document,
        "Розробку поділено на три логічні commits: структура й модель даних; "
        "business logic, консольний інтерфейс і тести; README з інструкціями. "
        "Після третього commit робоче дерево було чистим.",
    )
    add_heading(document, "11.1. Результат git status", 2)
    add_picture(
        document,
        "git_status.png",
        "Рисунок 4 — Чистий стан working tree після програмної частини",
        "Git status output showing a clean working tree on main",
    )
    add_heading(document, "11.2. Результат git log --oneline", 2)
    add_picture(
        document,
        "git_log.png",
        "Рисунок 5 — Три логічні commits програмної частини",
        "Git log with three logical commits",
    )

    add_heading(document, "12. Відповіді на контрольні запитання")
    answers = [
        (
            "Що таке Python-проєкт?",
            "Організована сукупність Python-коду, конфігурації, документації, "
            "тестів і ресурсів, які разом утворюють програмний продукт.",
        ),
        (
            "Чим проєкт відрізняється від одного Python script?",
            "Проєкт має модульну структуру, керування залежностями, "
            "документацію, тести й визначений спосіб встановлення та запуску.",
        ),
        ("Що таке Python module?", "Окремий файл .py, який містить Python-код."),
        (
            "Що таке Python package?",
            "Каталог, що об'єднує пов'язані Python modules і може бути "
            "імпортований як одна логічна одиниця.",
        ),
        (
            "Для чого потрібний __init__.py?",
            "Він позначає звичайний package, виконує його ініціалізацію та може "
            "визначати публічний API або метадані.",
        ),
        (
            "Що таке src-layout?",
            "Структура проєкту, у якій код встановлюваного пакета розміщено в "
            "окремому каталозі src.",
        ),
        (
            "Які переваги має src-layout?",
            "Він відокремлює код від службових файлів і змушує тести перевіряти "
            "встановлений пакет, а не випадковий імпорт із кореня репозиторію.",
        ),
        (
            "Для чого використовується virtual environment?",
            "Для ізоляції interpreter і залежностей конкретного проєкту від "
            "інших проєктів та системного Python.",
        ),
        ("Як створити .venv?", "Командою python -m venv .venv."),
        (
            "Як активувати virtual environment?",
            "Windows: .venv\\Scripts\\Activate.ps1; Linux/macOS: "
            "source .venv/bin/activate.",
        ),
        (
            "Для чого використовується pip?",
            "Для встановлення, оновлення й видалення Python packages та їхніх "
            "залежностей.",
        ),
        (
            "Чому рекомендується python -m pip?",
            "Такий запис запускає pip саме для обраного Python interpreter і "
            "зменшує ризик встановлення залежностей не в те середовище.",
        ),
        (
            "Для чого потрібний pyproject.toml?",
            "Для опису build system, метаданих, залежностей, console scripts і "
            "налаштувань інструментів розробки.",
        ),
        (
            "Що містить секція [project]?",
            "Назву, версію, опис, вимоги до Python, залежності та інші метадані "
            "пакета.",
        ),
        (
            "Що таке editable installation?",
            "Встановлення з параметром -e, за якого зміни в source code одразу "
            "доступні встановленому пакету без повторної інсталяції.",
        ),
        (
            "Для чого використовується if __name__ == \"__main__\"?",
            "Щоб application entry point виконувався лише за безпосереднього "
            "запуску модуля, але не під час імпорту.",
        ),
        ("Що таке Git?", "Розподілена система контролю версій."),
        (
            "Що таке repository?",
            "Сховище файлів проєкту разом із Git-історією, гілками й "
            "метаданими змін.",
        ),
        (
            "Що таке commit?",
            "Зафіксований стан підготовлених змін з унікальним ідентифікатором "
            "та повідомленням.",
        ),
        (
            "Для чого використовується git status?",
            "Для перегляду поточної гілки та стану working tree і staging area.",
        ),
        ("Для чого git add?", "Для додавання вибраних змін до staging area."),
        (
            "Для чого git commit?",
            "Для збереження підготовленого набору змін в історії repository.",
        ),
        ("Що показує git log?", "Історію commits та їхні метадані."),
        (
            "Для чого потрібен .gitignore?",
            "Щоб Git не відстежував virtual environment, кеші, build artifacts, "
            "налаштування IDE та секретні файли.",
        ),
        (
            "Чому .venv не потрібно комітити?",
            "Середовище велике, платформозалежне й відтворюється з "
            "pyproject.toml, тому його зберігання в repository недоцільне.",
        ),
        (
            "Для чого потрібний README?",
            "Для опису призначення, вимог, встановлення, запуску, тестування та "
            "структури проєкту.",
        ),
        (
            "Що таке type hints?",
            "Анотації очікуваних типів параметрів, повернених значень і змінних, "
            "які полегшують аналіз та супровід коду.",
        ),
        (
            "Для чого програму розділяють на modules?",
            "Для розподілу відповідальностей, повторного використання, "
            "незалежного тестування й простішого супроводу.",
        ),
        (
            "Що таке business logic?",
            "Правила й обчислення предметної області, які не залежать від "
            "конкретного інтерфейсу користувача.",
        ),
        (
            "Чому потрібно відокремлювати application entry point від business logic?",
            "Щоб логіку можна було тестувати й повторно використовувати без "
            "запуску консолі, а інтерфейс можна було змінювати окремо.",
        ),
    ]
    for number, (question, answer) in enumerate(answers, start=1):
        add_question_answer(document, number, question, answer)

    add_heading(document, "13. Висновок")
    add_body(
        document,
        "У лабораторній роботі створено встановлюваний Python-проєкт за "
        "src-layout для системи оцінювання студентів. Реалізовано всі функції "
        "варіанта №7, окрему модель даних, business logic, application entry "
        "point, type hints, валідацію, меню, розширений пошук, статистику й "
        "власний виняток. Пакет встановлено в .venv у editable mode. Усі 38 "
        "автоматизованих тестів пройшли, а coverage становить 99,6%. Код, "
        "README і конфігурацію зафіксовано трьома логічними Git commits.",
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    build_report()
