"""Build the laboratory work 2 report for variant 7."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_report import (
    add_body,
    add_code,
    add_cover_paragraph,
    add_picture,
    add_question_answer,
    clear_document_body,
    repeat_table_header,
    set_run_font,
)
from create_lab02_report_assets import render_terminal


ROOT = Path(__file__).resolve().parents[1]
SOURCE_TEMPLATE = ROOT / "docs" / "Lab_01_Report_Variant_7.docx"
OUTPUT = ROOT / "report" / "Звіт_ЛР02_варіант_7.docx"

sys.path.insert(0, str(ROOT / "src"))

from student_grading.benchmark import (  # noqa: E402
    BenchmarkResult,
    format_benchmark_table,
    run_benchmarks,
)


def configure_styles(document: Document) -> None:
    """Use the black Times New Roman hierarchy from the report reference."""

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")

    for style_name, size in (("Heading 1", 14), ("Heading 2", 13)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(6)


def set_cell_shading(cell, fill: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = cell_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        cell_properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), "D9D9D9")


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])

    for index, (header, width) in enumerate(zip(headers, widths, strict=True)):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "404040")
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, "Times New Roman", 10, bold=True)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        for column_index, (value, width) in enumerate(
            zip(values, widths, strict=True)
        ):
            cell = row.cells[column_index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "F2F2F2")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if column_index == 0
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            run = paragraph.add_run(value)
            set_run_font(run, "Times New Roman", 10)

    trailing = document.add_paragraph()
    trailing.paragraph_format.space_after = Pt(4)


def build_benchmark_figure(results: list[BenchmarkResult]) -> None:
    """Keep the report table and screenshot based on the same benchmark run."""

    output = format_benchmark_table(results)
    render_terminal(
        "student-grading — list vs dict benchmark",
        "$ python -m student_grading --benchmark\n" + output,
        ROOT / "docs" / "screenshots" / "lab02" / "benchmark_results.png",
    )


def build_report() -> None:
    if not SOURCE_TEMPLATE.exists():
        raise FileNotFoundError(f"Missing report template: {SOURCE_TEMPLATE}")

    benchmark_results = run_benchmarks()
    build_benchmark_figure(benchmark_results)

    OUTPUT.parent.mkdir(exist_ok=True)
    shutil.copy2(SOURCE_TEMPLATE, OUTPUT)
    document = Document(OUTPUT)
    clear_document_body(document)
    configure_styles(document)

    footer = document.sections[0].footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.clear()
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run(
        "Лабораторна робота №2 | Варіант №7 | Заставний Денис"
    )
    set_run_font(footer_run, "Times New Roman", 9)

    add_cover_paragraph(
        document,
        "ЛЬВІВСЬКИЙ НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ\nІМЕНІ ІВАНА ФРАНКА",
        bold=True,
        after=86,
    )
    add_cover_paragraph(
        document,
        "ЗВІТ\nДО ЛАБОРАТОРНОЇ РОБОТИ №2",
        size=16,
        bold=True,
        after=16,
    )
    add_cover_paragraph(
        document,
        "«Структури даних та функціональна декомпозиція»",
        bold=True,
        after=18,
    )
    add_cover_paragraph(
        document,
        "з курсу «Професійний Python»\n"
        "Варіант №7 — «Аналіз системи оцінювання студентів»",
        after=58,
    )
    add_cover_paragraph(
        document,
        "Виконав(ла): Заставний Денис",
        after=10,
        left_indent=3.15,
    )
    add_cover_paragraph(
        document,
        "Група: ФеП-32",
        after=10,
        left_indent=3.15,
    )
    add_cover_paragraph(
        document,
        "Перевірив(ла): Свелеба Сергій Андрійович",
        after=10,
        left_indent=3.15,
    )
    add_cover_paragraph(
        document,
        "Дата: 17.10.2026",
        after=58,
        left_indent=3.15,
    )
    add_cover_paragraph(document, "Львів — 2026")
    document.add_page_break()

    add_heading(document, "1. Тема та мета роботи")
    add_body(
        document,
        "Тема: «Структури даних та функціональна декомпозиція». Мета роботи — "
        "сформувати практичні навички професійного використання list, tuple, "
        "set, dict і структур модуля collections; навчитися обирати структуру "
        "відповідно до задачі, оцінювати складність операцій і декомпозувати "
        "програму на невеликі функції з чіткою відповідальністю.",
    )

    add_heading(document, "2. Індивідуальний варіант і постановка задачі")
    add_body(
        document,
        "Варіант №7 — «Аналіз системи оцінювання студентів». Для кожного "
        "запису зберігаються ID, ПІБ, дисципліна, tuple оцінок і група. "
        "Потрібно обчислити середній бал, мінімальну та максимальну оцінки, "
        "визначити успішність, відсортувати студентів, знайти записи вище "
        "порога, згрупувати дані за дисциплінами, використати Counter і "
        "створити dict-index за студентом або ID.",
    )
    add_body(
        document,
        "Додатково реалізовано всі ключові можливості підвищеної складності: "
        "універсальні filter і sort, параметризований decorator, closure, "
        "вкладене групування, Counter, deque історії операцій, benchmark, "
        "експорт JSON-підсумку та pipeline processing.",
    )

    add_heading(document, "3. Опис структури даних")
    add_body(
        document,
        "Наскрізна модель StudentRecord розширена полями group і student_id, "
        "не порушуючи функціональність лабораторної роботи №1. Об'єкт є "
        "незмінним dataclass зі slots; конструктор нормалізує текст, створює "
        "tuple оцінок і перевіряє діапазон 0–100 та додатність ID.",
    )
    add_code(
        document,
        "@dataclass(frozen=True, slots=True)\n"
        "class StudentRecord:\n"
        "    full_name: str\n"
        "    subject: str\n"
        "    grades: tuple[float, ...]\n"
        "    group: str\n"
        "    student_id: int | None",
    )

    add_heading(document, "4. Вибір list, tuple, set, dict і collections")
    add_table(
        document,
        ["Структура", "Застосування", "Причина вибору"],
        [
            ["list", "послідовність студентів і результати", "порядок, ітерація, сортування"],
            ["tuple", "оцінки та benchmark records", "незмінність структурованих значень"],
            ["set", "унікальні дисципліни", "унікальність і O(1) membership у середньому"],
            ["dict", "індекси, групи, агрегати", "пошук key у середньому O(1)"],
            ["Counter", "частоти дисциплін та оцінок", "готова частотна статистика"],
            ["defaultdict", "звичайне й вкладене групування", "автоматичне створення груп"],
            ["deque", "10 останніх операцій", "обмежена історія без ручного видалення"],
        ],
        [1.0, 2.2, 3.3],
    )

    add_heading(document, "5. Структура проєкту")
    add_code(
        document,
        "student-grading-system/\n"
        "├── pyproject.toml, README.md, .gitignore\n"
        "├── report/Звіт_ЛР02_варіант_7.docx\n"
        "├── src/student_grading/\n"
        "│   ├── models.py, data.py, services.py\n"
        "│   ├── processors.py, analytics.py\n"
        "│   ├── decorators.py, benchmark.py\n"
        "│   ├── config.py, exceptions.py\n"
        "│   └── main.py, __main__.py\n"
        "├── tests/\n"
        "│   ├── test_models.py, test_services.py, test_main.py\n"
        "│   └── test_processors.py, test_analytics.py,\n"
        "│       test_decorators.py, test_benchmark.py\n"
        "└── docs/screenshots/lab02/",
    )
    add_body(
        document,
        "data.py надає детермінований набір записів; processors.py містить "
        "операції над колекціями; analytics.py — агрегації, closure та export; "
        "decorators.py — параметризований decorator і deque; benchmark.py — "
        "відтворюваний експеримент; main.py — консольний entry point.",
    )

    add_heading(document, "6. Основні функції та декомпозиція")
    add_table(
        document,
        ["Функція", "Відповідальність"],
        [
            ["unique_subjects", "set comprehension унікальних дисциплін"],
            ["group_by_subject", "групування через defaultdict(list)"],
            ["count_subjects", "частоти дисциплін через Counter"],
            ["index_students_by_id", "dict-index для пошуку за ID"],
            ["filter_items / sort_items", "універсальні higher-order operations"],
            ["aggregate_by_subject", "агрегація статистики за дисциплінами"],
            ["create_average_filter", "closure із порогом середнього бала"],
            ["process_pipeline", "послідовне застосування *operations"],
            ["benchmark_search", "вимірювання list і dict lookup"],
            ["export_summary", "формування та запис JSON-підсумку"],
        ],
        [2.4, 4.1],
    )

    add_heading(document, "7. Comprehensions і generator expressions")
    add_body(
        document,
        "List comprehension формує відфільтровані результати, set "
        "comprehension — унікальні дисципліни, dict comprehension — індекс. "
        "Generator expression використано для Counter та агрегатів без "
        "створення зайвого проміжного list.",
    )
    add_code(
        document,
        "return [item for item in items if predicate(item)]\n\n"
        "return {student.subject for student in students}\n\n"
        "index = {\n"
        "    student.student_id: student\n"
        "    for student in records\n"
        "    if student.student_id is not None\n"
        "}\n\n"
        "Counter(student.subject for student in students)",
    )

    add_heading(document, "8. *args, **kwargs, lambda та функції вищого порядку")
    add_body(
        document,
        "Функція calculate_average_values(*values) приймає довільну кількість "
        "оцінок. create_student_record(**fields) створює валідований об'єкт із "
        "іменованих полів. Lambda передається у sort_items як key та в "
        "process_pipeline як компактна операція. Складна business logic "
        "залишається у звичайних іменованих функціях.",
    )
    add_code(
        document,
        "def calculate_average_values(*values: float) -> float:\n"
        "    return sum(values) / len(values) if values else 0.0\n\n"
        "def create_student_record(**fields: Any) -> StudentRecord:\n"
        "    return StudentRecord(**fields)\n\n"
        "sort_items(\n"
        "    students,\n"
        "    key=lambda student: student.average_grade,\n"
        "    reverse=True,\n"
        ")",
    )

    add_heading(document, "9. Closure")
    add_body(
        document,
        "create_average_filter(minimum_average) повертає predicate. Після "
        "завершення зовнішньої функції predicate пам'ятає нормалізований "
        "threshold і може багаторазово використовуватися для фільтрації.",
    )
    add_code(
        document,
        "def create_average_filter(minimum_average: float):\n"
        "    threshold = float(minimum_average)\n\n"
        "    def predicate(student: StudentRecord) -> bool:\n"
        "        return student.average_grade >= threshold\n\n"
        "    return predicate",
    )

    add_heading(document, "10. Параметризований decorator")
    add_body(
        document,
        "track_operation(label) є параметризованим decorator. Вкладені "
        "decorator і wrapper утворюють closures, functools.wraps зберігає "
        "метадані, а wrapper(*args, **kwargs) сумісний із різними signatures. "
        "Назва і тривалість додаються до deque(maxlen=10).",
    )
    add_code(
        document,
        "def track_operation(label: str | None = None):\n"
        "    def decorator(function):\n"
        "        operation_name = label or function.__name__\n\n"
        "        @wraps(function)\n"
        "        def wrapper(*args, **kwargs):\n"
        "            started = perf_counter()\n"
        "            try:\n"
        "                return function(*args, **kwargs)\n"
        "            finally:\n"
        "                history.append((operation_name, perf_counter() - started))\n"
        "        return wrapper\n"
        "    return decorator",
    )

    add_heading(document, "11. Сортування, фільтрація, групування, агрегація і пошук")
    add_body(
        document,
        "Сортування виконує sorted із key-функцією; фільтрація — універсальна "
        "filter_items; group_by_subject і group_by_subject_and_group будують "
        "звичайні та вкладені групи; aggregate_by_subject обчислює count, "
        "average, min і max; find_student_linear демонструє O(n), а готовий "
        "dict-index — O(1) у середньому. Вхідні списки не змінюються.",
    )

    add_heading(document, "12. Результат роботи програми")
    add_picture(
        document,
        "lab02/application_output_1.png",
        "Рисунок 1 — Вхідні дані, рейтинг і пошук за порогом",
        "Console output with student records, ranking, and threshold filtering",
        width=6.5,
    )
    add_picture(
        document,
        "lab02/application_output_2.png",
        "Рисунок 2 — Групування, Counter, індекс, агрегація та pipeline",
        "Console output with collection processing results",
        width=6.5,
    )

    add_heading(document, "13. Аналіз часової складності")
    add_table(
        document,
        ["Операція", "Структура", "Складність"],
        [
            ["Лінійний пошук ID", "list", "O(n)"],
            ["Побудова індексу", "dict comprehension", "O(n)"],
            ["Пошук ID у готовому індексі", "dict", "O(1) середнє"],
            ["Membership дисципліни", "set", "O(1) середнє"],
            ["Сортування студентів", "list / Timsort", "O(n log n)"],
            ["Групування й агрегація", "defaultdict / dict", "O(n)"],
        ],
        [3.0, 2.0, 1.5],
    )
    add_body(
        document,
        "Індекс доцільно будувати, коли після одноразових O(n) витрат потрібно "
        "виконати багато пошуків. Для одиничного пошуку малих даних list може "
        "бути достатнім, але при масштабуванні dict стабільно ефективніший.",
    )

    add_heading(document, "14. Експериментальна частина та benchmark")
    benchmark_rows = [
        [
            f"{result.record_count:,}".replace(",", " "),
            f"{result.list_search_seconds * 1_000_000:.3f}",
            f"{result.dict_search_seconds * 1_000_000:.3f}",
            f"{result.speedup:.1f}x",
        ]
        for result in benchmark_results
    ]
    add_table(
        document,
        ["Кількість records", "list search, мкс", "dict search, мкс", "Прискорення"],
        benchmark_rows,
        [1.6, 1.7, 1.7, 1.5],
    )
    add_picture(
        document,
        "lab02/benchmark_results.png",
        "Рисунок 3 — Фактичний benchmark list search проти dict search",
        "Terminal benchmark for 1000, 10000, and 100000 records",
        width=6.5,
    )
    add_body(
        document,
        "Експеримент шукає останній ID, тобто вимірює worst-case для list. "
        "Зі збільшенням набору в 100 разів час list search зростає приблизно "
        "лінійно, тоді як lookup у попередньо побудованому dict залишається "
        "майже сталим. Практичні результати узгоджуються з O(n) та O(1).",
    )

    add_heading(document, "15. Реалізовані завдання підвищеної складності")
    add_body(
        document,
        "Реалізовано десять додаткових можливостей із переліку: універсальний "
        "filter, універсальне sort, параметризований decorator, closure для "
        "filters, вкладене групування, статистика Counter, deque історії, "
        "benchmark list vs dict, export summary і pipeline processing. Це "
        "перевищує вимогу реалізувати щонайменше три можливості.",
    )

    add_heading(document, "16. Автоматизована перевірка")
    add_body(
        document,
        "Набір із 62 тестів перевіряє модель, усі операції Lab 01 і Lab 02, "
        "comprehensions, індекси, grouping, Counter, closure, decorator, deque, "
        "pipeline, benchmark, JSON export, CLI та помилки введення. Усі тести "
        "пройшли; branch coverage становить 99,8% при порозі 90%.",
    )
    add_picture(
        document,
        "lab02/test_results.png",
        "Рисунок 4 — Результат pytest і coverage",
        "Terminal output showing 62 passed tests and 99.8 percent coverage",
        width=6.5,
    )

    add_heading(document, "17. Відповіді на контрольні запитання")
    answers = [
        ("Що таке list?", "Впорядкована змінювана колекція Python."),
        ("Чим list відрізняється від tuple?", "List змінюваний, а tuple після створення незмінний."),
        ("Чому tuple називається immutable?", "Його елементи не можна додати, вилучити або замінити на місці."),
        ("Для чого використовується set?", "Для унікальних hashable елементів і швидких множинних операцій."),
        ("Яка властивість елементів set є важливою?", "Елементи мають бути hashable і мати стабільний hash."),
        ("Що таке dict?", "Відображення hashable keys на values."),
        ("Що таке key-value structure?", "Структура, де значення отримують за пов'язаним унікальним ключем."),
        ("Чому пошук key у dict зазвичай швидкий?", "Dict використовує hash table, тому середній lookup має O(1)."),
        ("Що таке hashability?", "Можливість об'єкта мати незмінний hash і коректно порівнюватися."),
        ("Чому list не може бути key у dict?", "List змінюваний і тому не є hashable."),
        ("Для чого використовується Counter?", "Для підрахунку частот hashable значень."),
        ("Для чого використовується defaultdict?", "Для автоматичного створення default value нового ключа."),
        ("Для чого використовується deque?", "Для O(1) додавання й видалення з обох кінців та bounded history."),
        ("Що таке list comprehension?", "Стислий спосіб побудови list з iterable і необов'язковою умовою."),
        ("Що таке set comprehension?", "Стислий спосіб побудови set унікальних значень."),
        ("Що таке dict comprehension?", "Стислий спосіб побудови пар key: value."),
        ("Що таке generator expression?", "Ліниве породження значень без створення повної проміжної collection."),
        ("Що означає O(1)?", "Кількість кроків у середньому не залежить від розміру input."),
        ("Що означає O(n)?", "Час зростає приблизно пропорційно кількості елементів."),
        ("Що означає O(n log n)?", "Типова складність ефективного comparison sorting."),
        ("Що означає O(n²)?", "При подвоєнні input робота зростає приблизно в чотири рази."),
        ("Яка приблизна складність пошуку у list?", "O(n) для пошуку довільного value."),
        ("Яка середня складність пошуку у set?", "O(1) у середньому."),
        ("Яка середня складність пошуку key у dict?", "O(1) у середньому."),
        ("Чому правильна структура даних може прискорити програму?", "Вона узгоджує часті операції з ефективним алгоритмом доступу."),
        ("Що таке функціональна декомпозиція?", "Поділ задачі на малі функції з однією відповідальністю."),
        ("Які переваги малих functions?", "Їх легше читати, тестувати, повторно використовувати й змінювати."),
        ("Що таке positional argument?", "Аргумент, значення якого визначається позицією у виклику."),
        ("Що таке keyword argument?", "Аргумент, переданий із явним ім'ям параметра."),
        ("Для чого використовується *args?", "Для приймання змінної кількості positional arguments як tuple."),
        ("Для чого використовується **kwargs?", "Для приймання змінної кількості keyword arguments як dict."),
        ("Що таке lambda function?", "Коротка анонімна функція з одним expression."),
        ("Де доцільно використовувати lambda?", "Для простих key, predicates і коротких callbacks."),
        ("Чому складну business logic не варто записувати в lambda?", "Це погіршує читабельність, тестування й діагностику."),
        ("Що таке closure?", "Внутрішня функція, що зберігає доступ до variables зовнішньої scope."),
        ("Які variables може пам'ятати closure?", "Referenced nonlocal variables, що існували під час його створення."),
        ("Що таке decorator?", "Callable, що обгортає інший callable та додає поведінку."),
        ("Для чого застосовуються decorators?", "Для logging, timing, validation, caching та інших cross-cutting concerns."),
        ("Для чого використовується functools.wraps?", "Щоб зберегти name, docstring та інші metadata обгорнутої функції."),
        ("Чому потрібно оцінювати складність алгоритму?", "Щоб прогнозувати масштабування й обирати відповідні structures та algorithms."),
    ]
    for number, (question, answer) in enumerate(answers, 1):
        add_question_answer(document, number, question, answer)

    add_heading(document, "18. Висновок")
    add_body(
        document,
        "У лабораторній роботі №2 для варіанта №7 створено модуль обробки "
        "структурованих даних системи оцінювання студентів. Свідомо застосовано "
        "list, tuple, set, dict, Counter, defaultdict і deque; реалізовано всі "
        "види comprehensions, generator expressions, *args, **kwargs, lambda, "
        "closure, параметризований decorator і функції вищого порядку. Програма "
        "виконує сортування, фільтрацію, групування, агрегацію та пошук, а "
        "benchmark на 1 000, 10 000 і 100 000 записів підтвердив різницю між "
        "O(n) list search і O(1) average dict lookup. Усі 62 тести пройшли з "
        "покриттям 99,8%, тому вимоги лабораторної роботи виконані повністю.",
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    build_report()
